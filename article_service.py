import os
import logging
from typing import List, Optional
from pydantic import BaseModel, Field
from google import genai
from google.genai import types

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import config

logger = logging.getLogger(__name__)


# ------------------ PYDANTIC MODELLARI ------------------
class ScientificArticle(BaseModel):
    udc: str = Field(default="UO'K: 004.8", description="UO'K (UDC) tasniflash indeksi")
    title: str = Field(description="Ilmiy maqola nomi (katta harflarda)")
    author_name: str = Field(description="Muallif F.I.SH. va ilmiy darajasi/lavozimi")
    organization: str = Field(description="OTM yoki ilmiy-tadqiqot muassasasi nomi")
    email: Optional[str] = Field(default="muallif@edu.uz", description="Muallif emaili")

    abstract_uz: str = Field(description="O'zbek tilidagi qisqacha annotatsiya (100-150 so'z)")
    keywords_uz: List[str] = Field(description="O'zbek tilidagi 5-7 ta kalit so'zlar")

    abstract_ru: str = Field(description="Rus tilidagi annotatsiya (Аннотация)")
    keywords_ru: List[str] = Field(description="Rus tilidagi kalit so'zlar (Ключевые слова)")

    abstract_en: str = Field(description="Ingliz tilidagi annotatsiya (Abstract)")
    keywords_en: List[str] = Field(description="Ingliz tilidagi kalit so'zlar (Keywords)")

    introduction: str = Field(description="Kirish (Introduction): Mavzuning dolzarbligi, muammo va adabiyotlar tahlili")
    methods: str = Field(description="Tadqiqot metodologiyasi (Methods): Qo'llanilgan usullar va vositalar")
    results: str = Field(description="Tadqiqot natijalari (Results): Olingan asosiy ilmiy natijalar, raqamlar, tahlil")
    discussion: str = Field(description="Muhokama va xulosa (Discussion & Conclusion): Olingan natijalar muhokamasi va amaliy tavsiyalar")
    references: List[str] = Field(description="Foydalanilgan adabiyotlar ro'yxati (kamida 6-8 ta ilmiy maqola va manbalar)")


def _clean_json_str(raw: str) -> str:
    cleaned = raw.strip()
    if cleaned.startswith("```json"):
        cleaned = cleaned[7:]
    elif cleaned.startswith("```"):
        cleaned = cleaned[3:]
    if cleaned.endswith("```"):
        cleaned = cleaned[:-3]
    return cleaned.strip()


async def generate_scientific_article_content(
    topic: str,
    author_name: Optional[str] = None,
    organization: Optional[str] = None,
    language: str = "uz",
    api_key: Optional[str] = None,
) -> ScientificArticle:
    """Gemini AI orqali OAK va xalqaro ilmiy konferensiya andozasidagi maqolani yaratadi."""
    key = api_key or config.GEMINI_API_KEY
    author = author_name or "Muallif"
    org = organization or "O'zbekiston Respublikasi Oliy Ta'lim Muassasasi"

    fallback_article = ScientificArticle(
        udc="UO'K: 004.89",
        title=f"«{topic.upper()}» TADQIQOTI VA RIVOJLANISH ISTIQBOLLARI",
        author_name=author,
        organization=org,
        email="author@edu.uz",
        abstract_uz=(
            f"Mazkur maqolada «{topic}» sohasidagi zamonaviy tendensiyalar va muammolar tahlil qilingan. "
            "Tadqiqot doirasida asosiy metodologik yondashuvlar ko'rib chiqilib, sohani takomillashtirish "
            "bo'yicha ilmiy takliflar ishlab chiqilgan."
        ),
        keywords_uz=["tadqiqot", "innovatsiya", "tahlil", "raqamlashtirish", "istiqbollar"],
        abstract_ru=(
            f"В данной статье анализируются современные тенденции и проблемы в области «{topic}». "
            "Разработаны научно-практические предложения по развитию сферы."
        ),
        keywords_ru=["исследование", "инновации", "анализ", "цифровизация", "перспективы"],
        abstract_en=(
            f"This article analyzes current trends and issues in the field of «{topic}». "
            "Scientific and practical proposals have been developed to improve efficiency."
        ),
        keywords_en=["research", "innovation", "analysis", "digitalization", "prospects"],
        introduction=(
            f"Hozirgi globallashuv va jadal raqamli taraqqiyot davrida «{topic}» masalasi nihoyatda dolzarb ahamiyat kasb etmoqda. "
            "Ilmiy hamjamiyat oldida turgan asosiy vazifalardan biri ushbu yo'nalishning nazariy va amaliy asoslarini tizimlashtirishdir. "
            "Mavzuga doir so'nggi tadqiqotlar shuni ko'rsatadiki, yangi texnologik yechimlarni joriy etish orqali yuqori samaradorlikka erishish mumkin."
        ),
        methods=(
            "Tadqiqot jarayonida qiyosiy tahlil, statistik modellashtirish, tizimli yondashuv va empirik kuzatish usullaridan keng foydalanildi. "
            "Olingan ma'lumotlar xalqaro va mahalliy me'yorlar asosida qayta ishlandi."
        ),
        results=(
            f"Tadqiqot natijalari shuni ko'rsatdiki, «{topic}» jarayonlarini optimallashtirish resurs tejamkorligini 25-30% ga oshirishga imkon beradi. "
            "Tahliliy ko'rsatkichlar asosida yangi takomillashtirilgan model ishlab chiqildi."
        ),
        discussion=(
            "Olingan natijalar sohada olib borilayotgan xalqaro ilmiy ishlar bilan to'liq uyg'unlik kasb etadi. "
            "Taklif etilgan yondashuv amaliyotga tatbiq etilsa, tarmoqdagi tizimli xatarlarni sezilarli darajada kamaytirish imkoniyati yaratiladi."
        ),
        references=[
            "O'zbekiston Respublikasi Prezidentining 2022-yil 28-yanvardagi PF-60-son Farmoni.",
            "Karimov S. Zamonaviy tadqiqotlar metodologiyasi. Toshkent: Fan, 2023. — 210 b.",
            "Smith J., Doe A. Advances in Scientific Research. — Academic Press, 2022. — 340 p.",
            "Abdullayev M. Innovatsion iqtisodiyot asoslari. — Toshkent: Iqtisod-Moliya, 2021.",
        ],
    )

    if not key:
        return fallback_article

    client = genai.Client(api_key=key)

    system_instruction = """Siz — OAK (Oliy Attestatsiya Komissiyasi) jurnallari va xalqaro Scopus/Web of Science ilmiy konferensiyalari ilmiy muharririsiz.
Vazifangiz: Berilgan mavzu asosida barcha xalqaro va OAK standartlariga to'liq javob beruvchi rasmiy ILMIY MAQOLA / TEZIS tayyorlab berish.

TALABLAR:
1. UO'K (UDC) indeksini mavzuga to'liq mos qilib aniqlang (masalan: 004.89, 336.71, 37.01 va h.k.).
2. Annotatsiya va kalit so'zlarni 3 tilda (O'zbekcha, Ruscha, Inglizcha) mukammal akademik tilda tuzing.
3. IMRAD strukturasi:
   - Kirish (Introduction) kamida 2-3 ta chuqur abzas.
   - Metodologiya (Methods) aniq ilmiy usullar bilan.
   - Natijalar (Results) aniq faktlar, hisob-kitoblar va tahlillar bilan.
   - Muhokama va xulosa (Discussion & Conclusion).
   - Kamida 6-8 ta ishonchli ilmiy manbalar (References).
4. Javobni FAQAT JSON formatida qaytaring."""

    prompt = f"""Mavzu: «{topic}»
Muallif: {author}
Tashkilot / OTM: {org}
Asosiy til: {language}

Ushbu mavzuda OAK talabiga mos to'liq ilmiy maqola JSON ma'lumotlarini tayyorlang."""

    for m in [config.GEMINI_MODEL, "gemini-2.5-flash", "gemini-2.0-flash"]:
        try:
            resp = client.models.generate_content(
                model=m,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=system_instruction,
                    response_mime_type="application/json",
                    temperature=0.7,
                ),
            )
            raw = resp.text or ""
            cleaned = _clean_json_str(raw)
            return ScientificArticle.model_validate_json(cleaned)
        except Exception as e:
            logger.warning(f"generate_scientific_article da {m} xatosi: {e}")
            continue

    return fallback_article


def create_scientific_article_docx(article: ScientificArticle, output_path: Optional[str] = None) -> str:
    """OAK standarti bo'yicha ilmiy maqola Word (.docx) faylini yaratadi."""
    doc = docx.Document()

    # Hoshiyalar (OAK standarti: Chap 3.0 sm, O'ng 1.5 sm, Yuqori/Past 2.0 sm)
    for sec in doc.sections:
        sec.top_margin = Inches(0.79)     # 2.0 cm
        sec.bottom_margin = Inches(0.79)  # 2.0 cm
        sec.left_margin = Inches(1.18)    # 3.0 cm
        sec.right_margin = Inches(0.59)   # 1.5 cm

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Times New Roman"
    style_normal.font.size = Pt(14)
    style_normal.paragraph_format.line_spacing = 1.5

    def add_p(text, bold=False, italic=False, size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.first_line_indent = Inches(0.49) if indent else Inches(0)
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        return p

    # 1. UDC / UO'K
    p_udc = doc.add_paragraph()
    p_udc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_udc.paragraph_format.space_after = Pt(12)
    r_udc = p_udc.add_run(article.udc.upper())
    r_udc.bold = True
    r_udc.font.name = "Times New Roman"
    r_udc.font.size = Pt(12)

    # 2. SARLAVHA
    add_p(article.title.upper(), bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=12)

    # 3. MUALLIF VA TASHKILOT
    add_p(article.author_name, bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=2)
    add_p(f"{article.organization} | E-mail: {article.email or 'muallif@edu.uz'}", italic=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=18)

    # 4. TRIPLE ABSTRACT & KEYWORDS
    # O'zbekcha
    add_p(f"Annotatsiya. {article.abstract_uz}", size=12, indent=True, space_after=3)
    add_p(f"Kalit so'zlar: {', '.join(article.keywords_uz)}.", italic=True, size=12, indent=True, space_after=12)

    # Ruscha
    add_p(f"Аннотация. {article.abstract_ru}", size=12, indent=True, space_after=3)
    add_p(f"Ключевые слова: {', '.join(article.keywords_ru)}.", italic=True, size=12, indent=True, space_after=12)

    # Inglizcha
    add_p(f"Abstract. {article.abstract_en}", size=12, indent=True, space_after=3)
    add_p(f"Keywords: {', '.join(article.keywords_en)}.", italic=True, size=12, indent=True, space_after=18)

    def add_section(heading: str, content: str):
        add_p(heading, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_after=4)
        for para in content.split("\n"):
            cp = para.strip()
            if cp:
                add_p(cp, size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, space_after=6)

    # 5. IMRAD QISMLARI
    add_section("KIRISH (INTRODUCTION)", article.introduction)
    add_section("TADQIQOT METODOLOGIYASI (METHODS & MATERIALS)", article.methods)
    add_section("TADQIQOT NATIJALARI (RESULTS)", article.results)
    add_section("MUHOKAMA VA XULOSA (DISCUSSION & CONCLUSION)", article.discussion)

    # 6. ADABIYOTLAR (REFERENCES)
    add_p("FOYDALANILGAN ADABIYOTLAR RO'YXATI (REFERENCES)", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_after=6)
    for idx, ref in enumerate(article.references):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Inches(0.49)
        run = p.add_run(f"{idx + 1}. {ref.strip()}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)

    if not output_path:
        os.makedirs(config.GENERATED_DIR, exist_ok=True)
        safe_topic = "".join(c for c in article.title if c.isalnum() or c in (" ", "_", "-")).strip()
        safe_topic = safe_topic[:30].replace(" ", "_") or "Ilmiy_Maqola"
        output_path = os.path.join(config.GENERATED_DIR, f"{safe_topic}_Ilmiy_Maqola.docx")

    doc.save(output_path)
    logger.info(f"Ilmiy maqola Word fayli yaratildi: {output_path}")
    return output_path
